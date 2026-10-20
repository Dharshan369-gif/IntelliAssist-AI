import streamlit as st
import os
import pandas as pd
import numpy as np
from dotenv import load_dotenv
from PyPDF2 import PdfReader
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from sklearn.metrics.pairwise import cosine_similarity
from google import genai

load_dotenv()
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

def extract_documents_with_metadata(uploaded_files):
    docs = []
    for file in uploaded_files:
        content = ""
        if file.name.endswith(".pdf"):
            reader = PdfReader(file)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    content += page_text + "\n"
        elif file.name.endswith(".docx"):
            doc_file = docx.Document(file)
            for para in doc_file.paragraphs:
                content += para.text + "\n"
        elif file.name.endswith(".txt"):
            content += file.read().decode("utf-8") + "\n"

        if content.strip():
            docs.append(Document(page_content=content, metadata={"source": file.name}))
    return docs

def get_text_chunks_and_df(docs):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    split_docs = text_splitter.split_documents(docs)

    df = pd.DataFrame([
        {
            "chunk_id": i + 1,
            "source": doc.metadata.get("source", "Unknown"),
            "char_length": len(doc.page_content),
            "content": doc.page_content
        }
        for i, doc in enumerate(split_docs)
    ])
    return split_docs, df

def get_vector_store(split_docs):
    embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
    vector_store = FAISS.from_documents(split_docs, embedding=embeddings)
    vector_store.save_local("faiss_index")

st.set_page_config(page_title="IntelliAssist AI", page_icon="📄", layout="wide")
st.title("📄 IntelliAssist AI: Smart Document Assistant")

with st.sidebar:
    st.header("Upload Documents")
    uploaded_files = st.file_uploader("Choose files", type=["pdf", "docx", "txt"], accept_multiple_files=True)

    if st.button("Process Documents"):
        if uploaded_files:
            with st.spinner("Extracting text and building vector database..."):
                raw_docs = extract_documents_with_metadata(uploaded_files)
                split_docs, chunk_df = get_text_chunks_and_df(raw_docs)
                get_vector_store(split_docs)
                
                st.session_state["chunk_df"] = chunk_df
                st.success(f"Processed {len(split_docs)} chunks across {len(uploaded_files)} file(s).")
        else:
            st.warning("Please upload at least one document first.")

    if "chunk_df" in st.session_state:
        st.subheader("Chunk Inspector")
        st.dataframe(
            st.session_state["chunk_df"][["chunk_id", "source", "char_length"]],
            use_container_width=True,
            hide_index=True
        )

st.subheader("Chat with your Document")
user_query = st.text_input("Ask a question about your uploaded documents:")

if user_query:
    st.info(f"You asked: {user_query}")
    
    with st.spinner("Searching and ranking relevant chunks..."):
        try:
            embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
            new_db = FAISS.load_local("faiss_index", embeddings, allow_dangerous_deserialization=True)
            
            docs = new_db.similarity_search(user_query, k=5)
            
            query_embedding = np.array(embeddings.embed_query(user_query)).reshape(1, -1)
            doc_embeddings = np.array(embeddings.embed_documents([d.page_content for d in docs]))
            
            scores = cosine_similarity(query_embedding, doc_embeddings).flatten()
            
            ranked_df = pd.DataFrame({
                "source": [d.metadata.get("source", "Unknown") for d in docs],
                "score": np.round(scores * 100, 2),
                "content": [d.page_content for d in docs]
            }).sort_values(by="score", ascending=False)
            
            confident_docs = ranked_df[ranked_df["score"] >= 1.0]
            
            if confident_docs.empty:
                st.warning("No relevant context found in documents matching this query closely enough.")
            else:
                context = "\n\n".join(confident_docs["content"].tolist())
                
                prompt = f"""
                Answer the question as detailed as possible from the provided context. If the answer is not in
                the provided context just say, "The answer is not available in the context", do not guess.
                
                Context:
                {context}
                
                Question: 
                {user_query}
                
                Answer:
                """
                
                response = client.models.generate_content(
                    model="gemini-3.5-flash-lite",
                    contents=prompt
                )
                
                st.success("🤖 Answer:")
                st.write(response.text)
                
                with st.expander("View Ranked Citations (Scikit-Learn & Pandas Output)"):
                    st.dataframe(
                        confident_docs[["source", "score", "content"]],
                        column_config={
                            "score": st.column_config.ProgressColumn(
                                "Relevance Match",
                                help="Cosine similarity score calculated via scikit-learn",
                                format="%.1f%%",
                                min_value=0,
                                max_value=100,
                            ),
                        },
                        use_container_width=True,
                        hide_index=True
                    )
                    
        except Exception as e:
            st.error(f"Something went wrong. Ensure documents are processed first. (Error: {e})")